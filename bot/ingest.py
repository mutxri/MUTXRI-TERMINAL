#!/usr/bin/env python3
"""bot/ingest.py - getting a private company's statements into the same shape.

Listed issuers arrive pre-parsed in static_data/financials. Everything else -
a private company's management accounts, an audited PDF, a spreadsheet a founder
emailed - has to be read off whatever format it came in. This module turns those
into the canonical document that bot/statements.py analyses, so a private company
gets exactly the same ratios, trends and flags as a listed one.

Deterministic readers first (CSV, XLSX, JSON, and PDF tables via pdfplumber).
When a PDF is too messy for table extraction - scanned, multi-column, footnoted -
bot/analyst.py can read it with Claude instead. That is extraction, not analysis:
the model reads numbers off a page, then every ratio is recomputed here in Python.
The model never does the arithmetic, and every extracted figure keeps the source
label it came from so a human can audit it against the document.
"""
import csv, datetime as dt, io, json, os, re

from . import statements as S

PERIOD_RE = re.compile(r"(?:FY\s*)?((?:19|20)\d{2})(?:\s*/\s*\d{2,4})?", re.I)


def _looks_like_period(cell):
    s = str(cell or "").strip()
    return bool(s) and bool(PERIOD_RE.search(s)) and len(s) <= 24


def _period_label(cell):
    m = PERIOD_RE.search(str(cell or ""))
    return ("FY" + m.group(1)) if m else str(cell).strip()


def _table_to_sections(rows, entity, source, section_hint=None):
    """Turn a label-plus-values grid into a canonical document.

    The header row is the first row with two or more year-like cells; every row
    after it contributes a label and its values. Rows whose label maps to no
    canonical line item are kept in `unmapped` rather than discarded, so nothing
    silently vanishes from a statement.
    """
    header_idx, periods, period_cols = None, [], []
    for i, row in enumerate(rows[:25]):
        yrs = [(j, c) for j, c in enumerate(row) if _looks_like_period(c)]
        if len(yrs) >= 2:
            header_idx = i
            period_cols = [j for j, _ in yrs]
            periods = [_period_label(c) for _, c in yrs]
            break
    if header_idx is None:
        raise ValueError("could not find a period header row (need >=2 year columns)")

    # Newest-first, matching the parsed-filing convention used across the terminal.
    order = sorted(range(len(periods)),
                   key=lambda k: S._fy(periods[k]) or 0, reverse=True)
    periods = [periods[k] for k in order]
    period_cols = [period_cols[k] for k in order]

    sections = {"income": {}, "balance": {}, "cashflow": {}}
    unmapped, current = [], section_hint
    SECTION_CUES = [
        (r"income statement|profit (?:and|&) loss|statement of (?:comprehensive )?income|p&l", "income"),
        (r"balance sheet|statement of financial position", "balance"),
        (r"cash ?flow", "cashflow"),
    ]

    for row in rows[header_idx + 1:]:
        if not row:
            continue
        label = str(row[0] or "").strip()
        if not label:
            continue
        low = label.lower()
        hit_section = next((sec for pat, sec in SECTION_CUES if re.search(pat, low)), None)
        values = [S._num(row[j]) if j < len(row) else None for j in period_cols]
        if hit_section and all(v is None for v in values):
            current = hit_section          # a section heading, not a data row
            continue
        if all(v is None for v in values):
            continue
        # Try the hinted section first, then any section that recognises the label.
        placed = False
        for sec in ([current] if current else []) + ["income", "balance", "cashflow"]:
            if not sec:
                continue
            key = S._canon(label, S.SECTION_MAPS[sec])
            if key and key not in sections[sec]:
                sections[sec][key] = values
                placed = True
                break
        if not placed:
            unmapped.append({"label": label, "values": values})

    if not any(sections.values()):
        raise ValueError("no recognisable statement lines found")

    return {
        "entity": entity,
        "periods": periods,
        "sections": {k: v for k, v in sections.items() if v},
        "source": source,
        "unmapped": unmapped,
    }


def _entity(name, kind="private", **kw):
    e = {"id": kw.get("id") or re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").upper(),
         "name": name, "kind": kind, "ticker": kw.get("ticker"),
         "currency": kw.get("currency"), "sector": kw.get("sector"),
         "country": kw.get("country")}
    return e


# ------------------------------------------------------------------- readers
def from_csv(path, name=None, **kw):
    """Read a CSV/TSV grid: first column labels, later columns periods."""
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        sample = f.read(8192)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [r for r in csv.reader(f, dialect)]
    return _table_to_sections(
        rows, _entity(name or os.path.splitext(os.path.basename(path))[0], **kw),
        {"kind": "csv", "detail": os.path.basename(path),
         "ingestedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")})


def from_xlsx(path, name=None, sheet=None, **kw):
    """Read the first (or named) worksheet as a label/period grid."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is required to read .xlsx - pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    return _table_to_sections(
        rows, _entity(name or os.path.splitext(os.path.basename(path))[0], **kw),
        {"kind": "xlsx", "detail": "%s#%s" % (os.path.basename(path), ws.title),
         "ingestedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")})


def from_json(path, name=None, **kw):
    """Read a statement already in the terminal's {rows,periods} shape."""
    raw = json.load(open(path, encoding="utf-8"))
    if isinstance(raw, dict) and "sections" in raw and "periods" in raw:
        return raw                                   # already canonical
    sections, periods = {}, raw.get("periods") or []
    for section in ("income", "balance", "cashflow"):
        blob = raw.get(section)
        if blob:
            vals, _ = S.normalise_section(blob if isinstance(blob, dict) else {"rows": blob},
                                          section)
            if vals:
                sections[section] = vals
    if not sections and raw.get("rows"):
        guess = raw.get("statement") or "income"
        vals, _ = S.normalise_section(raw, guess)
        sections[guess] = vals
    if not sections:
        raise ValueError("unrecognised JSON statement shape")
    return {"entity": _entity(name or raw.get("name") or "entity", **kw),
            "periods": periods, "sections": sections,
            "source": {"kind": "json", "detail": os.path.basename(path)}}


def pdf_tables(path, max_pages=40):
    """Extract candidate tables from a text-based PDF. Returns list of grids."""
    try:
        import pdfplumber
    except ImportError:
        raise RuntimeError("pdfplumber is required to read PDFs - pip install pdfplumber")
    grids = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:max_pages]:
            for tbl in page.extract_tables() or []:
                if tbl and len(tbl) >= 3:
                    grids.append(tbl)
    return grids


def pdf_text(path, max_pages=40):
    """Plain text of a PDF, for the model-assisted path and for auditing."""
    try:
        import pdfplumber
    except ImportError:
        raise RuntimeError("pdfplumber is required to read PDFs - pip install pdfplumber")
    out = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:max_pages]:
            out.append(page.extract_text() or "")
    return "\n".join(out)


# One figure. Thousands are comma-grouped ("35,946") or space-grouped in the
# South African style ("7 535"); a space is only part of a number when it is
# followed by exactly three digits, otherwise it separates two columns. Getting
# this wrong fuses "35,946 41,083" into 3594641083 - two years of revenue read
# as one impossible number.
_NUM_TOKEN = re.compile(r"\(?-?\d{1,3}(?:,\d{3})+(?:\.\d+)?\)?"
                        r"|\(?-?\d{1,3}(?:\s\d{3})+(?:\.\d+)?\)?"
                        r"|\(?-?\d+(?:\.\d+)?\)?%?")


def _parse_statement_line(line):
    """Split one text line into (label, values) pairs.

    Filings print two statements side by side, so a single extracted line can
    read "Profit after tax 5,246 4,483 Non-current assets 9,687 10,061" - the
    income statement and the balance sheet on the same row of the page. Scanning
    label-then-figures-then-label recovers both instead of mangling the first
    label and losing the second entirely.
    """
    out, i, n = [], 0, len(line)
    while i < n:
        m = _NUM_TOKEN.search(line, i)
        if not m:
            break
        label = line[i:m.start()].strip(" .: ")
        vals, j = [], m.start()
        while True:
            m2 = _NUM_TOKEN.match(line, j)
            if not m2:
                break
            tok = m2.group(0)
            if tok.endswith("%"):        # a percentage is commentary, not a column
                j = m2.end()
                while j < n and line[j] == " ":
                    j += 1
                continue
            v = S._num(tok)
            if v is not None:
                vals.append(v)
            j = m2.end()
            while j < n and line[j] == " ":
                j += 1
        if label and vals and re.search(r"[A-Za-z]", label) and len(label) >= 3:
            out.append({"label": label, "values": vals})
        i = j if j > i else m.end()
    return out
# Period headers as filings actually write them: "31-Dec-25", "31 December 2025",
# "FY2025", "2025". Two-digit years appear in Kenyan and Nigerian releases.
_HDR_TOKEN = re.compile(
    r"(?:\d{1,2}[-\s](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*[-\s]"
    r"(?P<yy>\d{2,4})|(?:FY\s*)?(?P<yyyy>(?:19|20)\d{2}))", re.I)


def _header_years(line):
    """Fiscal years named on one line, newest-first order preserved as written."""
    out = []
    for m in _HDR_TOKEN.finditer(line):
        y = m.group("yyyy") or m.group("yy")
        if not y:
            continue
        y = int(y)
        if y < 100:                      # "31-Dec-25" -> 2025
            y += 2000
        if 1990 <= y <= 2100 and ("FY%d" % y) not in out:
            out.append("FY%d" % y)
    return out


def pdf_line_items(path, max_pages=40):
    """Read statement lines straight out of the PDF's text layer.

    Real filings lay their statements out with whitespace rather than ruled
    cells, so pdfplumber's table extractor returns value-only columns and drops
    the labels entirely - which makes the table path useless on every real
    document tested here. The text layer keeps them: "Net revenue 23,192 25,716"
    is one line, and a label followed by figures is all a statement line is.

    Two-column pages interleave the income statement with the balance sheet, but
    that is harmless: each line is mapped on its own label, so the order the
    lines arrive in does not matter.
    """
    text = pdf_text(path, max_pages=max_pages)
    periods, rows = [], []
    for raw in text.split("\n"):
        line = raw.replace("\u00a0", " ").strip()
        if not line:
            continue
        yrs = _header_years(line)
        # A period header names years and is mostly dates. A statement line that
        # merely mentions a year ("Dividend for 2025  1,200  900") is not one.
        if len(yrs) >= 2 and len(re.sub(r"[^A-Za-z]", "", line)) < 40:
            if len(yrs) > len(periods):
                periods = yrs
            continue
        rows.extend(_parse_statement_line(line))
    return periods, rows


def _lines_to_sections(periods, rows, entity, source):
    """Build the canonical document from label/value lines."""
    if len(periods) < 1:
        raise ValueError("no period header found in the text layer")
    sections = {"income": {}, "balance": {}, "cashflow": {}}
    unmapped = []
    n = len(periods)
    for r in rows:
        vals = r["values"]
        if len(vals) < n:
            # A line reporting fewer figures than there are periods cannot be
            # aligned to them safely; keep it visible rather than guess.
            unmapped.append(r["label"])
            continue
        # Statements print "Label | Note | FY2025 | FY2024", so a row with more
        # figures than periods usually carries a note reference on the left.
        # Taking the trailing values keeps the money columns; taking the leading
        # ones read Aveng's revenue as 27 and ArcelorMittal's as 4.
        vals = vals[-n:]
        placed = False
        for sec, mapping in (("income", S.INCOME_MAP), ("balance", S.BALANCE_MAP),
                             ("cashflow", S.CASHFLOW_MAP)):
            key = S._canon(r["label"], mapping)
            if key and key not in sections[sec]:
                sections[sec][key] = vals
                placed = True
                break
        if not placed:
            unmapped.append(r["label"])
    if not any(sections.values()):
        raise ValueError("no recognisable statement lines in the text layer")
    dropped = _drop_scale_outliers(sections)
    if not any(sections.values()):
        raise ValueError("no statement lines survived the scale check")
    return {"entity": entity, "periods": periods, "sections": sections,
            "source": source, "unmapped": unmapped[:60],
            "droppedOutOfScale": dropped}


# Line items whose value is legitimately small next to the rest of a statement,
# so they must be exempt from the scale check.
_SMALL_BY_NATURE = {"eps"}


def _drop_scale_outliers(sections):
    """Discard figures that cannot belong to the same statement.

    A note reference left stranded on its own line ("Revenue 27") parses as a
    perfectly good figure. Nothing about the line says it is wrong - only its
    size relative to the rest of the document does, and a revenue of 27 sitting
    beside a profit of 2,717,172 is a note number, not money. Anything more than
    three orders of magnitude below the document's median figure is dropped and
    reported rather than published as a company's revenue.
    """
    mags = []
    for sec, items in sections.items():
        for key, vals in items.items():
            if key in _SMALL_BY_NATURE:
                continue
            mags += [abs(v) for v in vals if v]
    if len(mags) < 4:
        return []
    mags.sort()
    median = mags[len(mags) // 2]
    floor = median / 1000.0
    dropped = []
    for sec, items in sections.items():
        for key in list(items):
            if key in _SMALL_BY_NATURE:
                continue
            vals = [v for v in items[key] if v]
            if vals and max(abs(v) for v in vals) < floor:
                dropped.append({"section": sec, "item": key, "values": items[key],
                                "medianFigure": median})
                del items[key]
    return dropped


def from_pdf(path, name=None, **kw):
    """Deterministic PDF path.

    Tries the text layer first - it is what actually works on real filings -
    then falls back to extracted tables for documents that genuinely are ruled
    grids. Raises when neither yields a statement, so the caller can fall back
    to the model-assisted reader rather than guessing here.
    """
    entity = _entity(name or os.path.splitext(os.path.basename(path))[0], **kw)
    stamp = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")

    try:
        periods, rows = pdf_line_items(path)
        if periods and rows:
            return _lines_to_sections(
                periods, rows, entity,
                {"kind": "pdf-text", "detail": os.path.basename(path),
                 "ingestedAt": stamp, "linesRead": len(rows)})
    except ValueError:
        pass

    grids = pdf_tables(path)
    for g in grids:
        try:
            return _table_to_sections(
                g, entity,
                {"kind": "pdf-table", "detail": os.path.basename(path),
                 "ingestedAt": stamp})
        except ValueError:
            continue
    raise ValueError("no parsable statement in %s (text layer and %d tables tried)"
                     % (os.path.basename(path), len(grids)))


READERS = {".csv": from_csv, ".tsv": from_csv, ".txt": from_csv,
           ".xlsx": from_xlsx, ".xlsm": from_xlsx,
           ".json": from_json, ".pdf": from_pdf}


def ingest(path, name=None, **kw):
    """Read any supported statement file into the canonical document shape."""
    ext = os.path.splitext(path)[1].lower()
    reader = READERS.get(ext)
    if not reader:
        raise ValueError("unsupported statement format %r (support: %s)"
                         % (ext, ", ".join(sorted(READERS))))
    return reader(path, name=name, **kw)


# ------------------------------------------------------------------- storage
PRIVATE_DIR = os.path.join(S.SD, "private")


def save(doc):
    """Persist an ingested private company so later runs can reuse it."""
    os.makedirs(PRIVATE_DIR, exist_ok=True)
    eid = doc["entity"]["id"]
    p = os.path.join(PRIVATE_DIR, "%s.json" % re.sub(r"[^A-Za-z0-9_.-]", "_", eid))
    with open(p, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
    return p


def load_private(entity_id):
    p = os.path.join(PRIVATE_DIR, "%s.json"
                     % re.sub(r"[^A-Za-z0-9_.-]", "_", entity_id))
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else None


def available_private():
    if not os.path.isdir(PRIVATE_DIR):
        return []
    return sorted(os.path.splitext(f)[0] for f in os.listdir(PRIVATE_DIR)
                  if f.endswith(".json"))


def load_any(entity_id):
    """Resolve an id to a document, public statements first, then private."""
    return S.load_public(entity_id) or load_private(entity_id)


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print(__doc__)
        print("usage: python -m bot.ingest <file> [--name NAME] [--save]")
        raise SystemExit(1)
    src = sys.argv[1]
    nm = None
    if "--name" in sys.argv:
        nm = sys.argv[sys.argv.index("--name") + 1]
    d = ingest(src, name=nm)
    a = S.analyse(d)
    print("%s (%s) - periods %s" % (a["entity"]["name"], a["entity"]["kind"],
                                    ", ".join(a["periods"])))
    for r in a["metrics"]:
        print("  %-8s rev %14s net %13s margin %7s" %
              (r["period"],
               "%.0f" % r["revenue"] if r["revenue"] is not None else "-",
               "%.0f" % r["net_profit"] if r["net_profit"] is not None else "-",
               "%.1f%%" % r["net_margin"] if r["net_margin"] is not None else "-"))
    for f in a["flags"]:
        print("  [%s] %s - %s" % (f["severity"].upper(), f["label"], f["detail"]))
    if d.get("unmapped"):
        print("  unmapped lines kept:", len(d["unmapped"]))
    if "--save" in sys.argv:
        print("saved:", save(d))
